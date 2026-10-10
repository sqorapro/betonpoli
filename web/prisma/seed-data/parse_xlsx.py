"""Разовый парсер Excel от заказчика -> JSON для сида.

python prisma/seed-data/parse_xlsx.py ../КАЛЬКУЛЯТОРЫ.xlsx "../Примеры товаров.xlsx"
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).parent


def num(v):
    if isinstance(v, (int, float)):
        return float(v)
    return None


def parse_products(path):
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    out = []
    for r in rows:
        if not r[2]:
            continue
        _, cat, name, sku, unit, price, chars, note = r[:8]
        m = re.search(r"([\d\s]+[,.]?\d*)", str(price).replace("\xa0", " "))
        price_val = float(m.group(1).replace(" ", "").replace(",", ".")) if m else None
        out.append({
            "category": re.sub(r"^\d+\.\s*", "", cat).strip(),
            "name": name.strip(),
            "sku": str(sku).strip(),
            "unit": unit.strip(),
            "priceFrom": price_val,
            "priceRaw": str(price),
            "characteristics": (chars or "").strip(),
            "description": (note or "").strip(),
        })
    return out


def parse_kits(path):
    wb = openpyxl.load_workbook(path)
    kits = []
    for ws in wb:
        # блоки: колонка с «Наименование» — B (обычно) или C (Lincrete)
        cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
        max_row = ws.max_row
        for off in (0, 1):  # смещение колонок: A..G или B..H
            col = lambda k: k + off  # noqa: E731 1=A
            r = 1
            while r <= max_row:
                v = cells.get((r, col(1)))
                if isinstance(v, str) and v.startswith("Площадь покрытия") or isinstance(v, str) and v.startswith("Длина плинтуса"):
                    kit = parse_block(cells, r, col, ws.title)
                    if kit:
                        kits.append(kit)
                r += 1
    return kits


def parse_block(cells, r, col, sheet):
    title = None
    for back in range(1, 3):
        t = cells.get((r - back, col(1)))
        if isinstance(t, str) and not t.startswith(("Итого", "Стоимость")):
            title = t.strip()
            break
    if not title:
        title = sheet
    uses_length = cells.get((r, col(1))).startswith("Длина")
    area = num(cells.get((r, col(6))))
    thickness = None
    rr = r + 1
    if isinstance(cells.get((rr, col(1))), str) and cells[(rr, col(1))].startswith("Толщина"):
        thickness = num(cells.get((rr, col(6))))
        rr += 1
    header = cells.get((rr, col(2)))
    if header != "Наименование":
        return None
    items = []
    rr += 1
    while isinstance(cells.get((rr, col(1))), (int, float)):
        items.append(parse_item(cells, rr, col))
        rr += 1
    if uses_length:
        for it in items:
            if it["type"] == "PER_AREA":
                it["type"] = "PER_LENGTH"
    return {"sheet": sheet, "title": title, "usesLength": uses_length,
            "defaultArea": area, "defaultThickness": thickness, "items": items}


def parse_item(cells, r, col):
    name = str(cells.get((r, col(2)))).strip()
    rate_raw = cells.get((r, col(3)))
    price = num(cells.get((r, col(4))))
    e = cells.get((r, col(5)))
    f = cells.get((r, col(6)))
    item = {"name": name, "pricePerUnit": price, "rateRaw": rate_raw, "eRaw": e, "fRaw": f}

    # объём тары: из E (=30*F4 / =F5*20) или из делителя в F (/30,1)
    pack = None
    if isinstance(e, str):
        m = re.fullmatch(r"=\s*([\d.]+)\s*\*\s*[A-Z]+\d+", e) or re.fullmatch(r"=\s*[A-Z]+\d+\s*\*\s*([\d.]+)", e)
        if m:
            pack = float(m.group(1))
    if pack is None and isinstance(f, str):
        m = re.search(r"/\s*([\d.]+)\s*,\s*1\)", f)
        if m:
            pack = float(m.group(1))
    if isinstance(e, str) and "CEILING.PRECISE" in e:  # лак: E = кг с округлением до 5
        m = re.search(r",\s*([\d.]+)\)", e)
        pack = float(m.group(1)) if m else pack

    mm = isinstance(rate_raw, str) and (
        re.fullmatch(r"=\s*[A-Z]+\d+\s*\*\s*([\d.]+)", rate_raw) or re.fullmatch(r"=\s*([\d.]+)\s*\*\s*[A-Z]+\d+", rate_raw))
    if mm:
        item.update(type="PER_AREA_MM", rate=float(mm.group(1)))
    elif num(rate_raw) is not None:
        item.update(type="PER_AREA", rate=num(rate_raw))
    elif isinstance(f, str) and re.search(r"CEILING\(\$?[A-Z]+\$?\d+/(\d+)", f):
        per = float(re.search(r"CEILING\(\$?[A-Z]+\$?\d+/(\d+)", f).group(1))
        # 1 рулон на `per` м²; E = кол-во × pack → это м² сетки
        item.update(type="PER_AREA", rate=1.0)
        pack = per
    elif num(f) is not None:
        item.update(type="FIXED", rate=num(f))
        if num(e) is not None and num(f):
            pack = num(e) / num(f)
    else:
        item.update(type="UNKNOWN", rate=None)
    if pack is None and num(e) is not None and num(f):
        pack = num(e) / num(f)  # кол-во задано руками: 180 м ленты / 6 шт = 30 м
    item["pack"] = pack
    return item


if __name__ == "__main__":
    calc, examples = sys.argv[1], sys.argv[2]
    (OUT / "products.json").write_text(json.dumps(parse_products(examples), ensure_ascii=False, indent=2), "utf-8")
    (OUT / "kits.raw.json").write_text(json.dumps(parse_kits(calc), ensure_ascii=False, indent=2, default=str), "utf-8")
    print("ok")
